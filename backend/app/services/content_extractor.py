import csv
import io
import logging
import os
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Dict, Any, Optional, List

from app.config import settings
from app.services.ocr_service import ocr_service
from app.services.vision_service import vision_service

logger = logging.getLogger("lazyflow.extractor")

# Maximum characters to extract per file to ensure minimal data transmission and strict privacy
MAX_SNIPPET_LENGTH = 1500
MAX_PREVIEW_ROWS = 4


class ContentExtractor:
    """
    Content extraction service for LazyFlow AI.
    Extracts readable text and structural metadata from files
    (PDF, DOCX, CSV, XLSX, Images, Code, Text) to enable content-aware
    purpose classification without sending full file binaries.
    """

    def extract_file_content(
        self,
        file_path: Path | str,
        filename: str,
        mime_type: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Extract readable content snippet and structural metadata from a file.
        Strictly preserves privacy: truncates snippets to minimal length needed
        for categorization, and never logs extracted content.
        """
        path = Path(file_path)
        if not path.exists() or not path.is_file():
            logger.info(f"[Diagnostic:Extraction] File '{filename}': not found on disk at {file_path}")
            return {
                "has_content": False,
                "content_snippet": "",
                "content_type": "missing",
                "extracted_metadata": {"error": "file_not_found", "is_empty_or_minimal": True}
            }

        file_size = path.stat().st_size
        if file_size == 0:
            logger.info(f"[Diagnostic:Extraction] File '{filename}': empty file (0 bytes)")
            return {
                "has_content": False,
                "content_snippet": "",
                "content_type": "empty",
                "extracted_metadata": {"is_empty_or_minimal": True, "size_bytes": 0}
            }

        ext = path.suffix.lower().lstrip(".")
        mime = (mime_type or "").lower()

        try:
            # 1. PDF Documents
            if ext == "pdf" or "pdf" in mime:
                return self._extract_pdf(path, filename)

            # 2. Word Documents (.docx)
            if ext == "docx" or "wordprocessingml" in mime:
                return self._extract_docx(path, filename)

            # 3. CSV Spreadsheets
            if ext == "csv" or "csv" in mime:
                return self._extract_csv(path, filename)

            # 4. Excel Spreadsheets (.xlsx)
            if ext == "xlsx" or "spreadsheetml" in mime:
                return self._extract_xlsx(path, filename)

            # 5. PowerPoint Presentations (.pptx)
            if ext == "pptx" or "presentationml" in mime:
                return self._extract_pptx(path, filename)

            # 6. Images (.png, .jpg, .jpeg, .webp, .bmp, .gif, .tiff)
            if ext in ("png", "jpg", "jpeg", "webp", "bmp", "gif", "tiff") or mime.startswith("image/"):
                return self._extract_image_metadata(path, filename)

            # 7. Text and Source Code (.txt, .md, .py, .js, .ts, .json, .html, .css, .yml, .sql, etc.)
            if ext in (
                "txt", "md", "markdown", "py", "js", "ts", "jsx", "tsx",
                "json", "html", "htm", "css", "scss", "yml", "yaml", "xml",
                "sql", "sh", "bat", "ps1", "c", "cpp", "h", "java", "rs", "go"
            ) or mime.startswith("text/") or "json" in mime:
                return self._extract_text_or_code(path, filename, ext)

            # 8. Unsupported or generic binary formats (archives, audio, video)
            return self._extract_generic_binary_info(path, filename, ext, mime)

        except Exception as e:
            logger.warning(f"[Diagnostic:Extraction] Safe extraction error for '{filename}': {type(e).__name__}: {e}")
            return {
                "has_content": False,
                "content_snippet": "",
                "content_type": "extraction_error",
                "extracted_metadata": {"error": str(type(e).__name__), "is_empty_or_minimal": True}
            }

    def _extract_pdf(self, path: Path, filename: str) -> Dict[str, Any]:
        """
        Extract readable text from PDF pages using PyMuPDF (fitz) and pypdf with fallback.
        Safe diagnostics log only metadata and character counts, never sensitive text.
        """
        text = ""
        page_count = 0
        is_encrypted = False
        method_used = "none"

        # Strategy 1: PyMuPDF (fitz) - fastest, handles complex font encodings and CID keys
        try:
            import fitz
            doc = fitz.open(str(path))
            page_count = len(doc)
            is_encrypted = bool(doc.is_encrypted)
            if not is_encrypted:
                extracted_pages = []
                for i in range(min(5, len(doc))):
                    page_text = doc[i].get_text() or ""
                    if page_text.strip():
                        extracted_pages.append(page_text.strip())
                text = "\n".join(extracted_pages)
                if text.strip():
                    method_used = "pymupdf_fitz"
            doc.close()
        except Exception as e:
            logger.debug(f"PyMuPDF fitz extraction attempt failed for {filename}: {e}")

        # Strategy 2: pypdf if fitz didn't yield text
        if not text.strip():
            try:
                from pypdf import PdfReader
                reader = PdfReader(str(path))
                is_encrypted = bool(reader.is_encrypted)
                if not is_encrypted:
                    page_count = len(reader.pages)
                    extracted_pages = []
                    for p in reader.pages[:5]:
                        t = p.extract_text() or ""
                        if t.strip():
                            extracted_pages.append(t.strip())
                    text = "\n".join(extracted_pages)
                    if text.strip():
                        method_used = "pypdf"
            except Exception as e:
                logger.debug(f"pypdf extraction attempt failed for {filename}: {e}")

        # Strategy 3: pdfplumber if still empty
        if not text.strip():
            try:
                import pdfplumber
                with pdfplumber.open(str(path)) as pdf:
                    page_count = len(pdf.pages)
                    extracted_pages = []
                    for page in pdf.pages[:5]:
                        t = page.extract_text() or ""
                        if t.strip():
                            extracted_pages.append(t.strip())
                    text = "\n".join(extracted_pages)
                    if text.strip():
                        method_used = "pdfplumber"
            except Exception as e:
                logger.debug(f"pdfplumber extraction attempt failed for {filename}: {e}")

        clean_text = self._clean_text(text)
        char_count = len(clean_text)
        has_content = char_count >= 20
        ocr_attempted = False
        ocr_used = False
        ocr_error = None

        # Strategy 4: OCR fallback for scanned or image-based PDFs when embedded text is empty or insufficient
        if not has_content and not is_encrypted:
            ocr_attempted = True
            if ocr_service.is_available():
                ocr_res = ocr_service.ocr_pdf(
                    path,
                    filename=filename,
                    max_pages=settings.OCR_MAX_PAGES
                )
                if ocr_res.get("success") and ocr_res.get("text"):
                    clean_ocr = self._clean_text(ocr_res["text"])
                    if len(clean_ocr) >= 15:
                        clean_text = clean_ocr
                        char_count = len(clean_text)
                        has_content = True
                        method_used = "ocr_tesseract"
                        ocr_used = True
                else:
                    ocr_error = ocr_res.get("error")
            else:
                ocr_error = "OCR engine not available"

        # Safe diagnostic logging (no text content logged!)
        logger.info(
            f"[Diagnostic:Extraction] PDF '{filename}' ({path.stat().st_size} bytes): "
            f"pages={page_count}, text_extracted={has_content}, char_count={char_count}, "
            f"parser='{method_used}', encrypted={is_encrypted}, ocr_used={ocr_used}"
        )

        return {
            "has_content": has_content,
            "content_snippet": clean_text[:MAX_SNIPPET_LENGTH] if has_content else "",
            "content_type": "pdf_scanned_ocr" if ocr_used else "pdf_text",
            "extracted_metadata": {
                "page_count": page_count,
                "is_encrypted": is_encrypted,
                "is_empty_or_minimal": not has_content,
                "char_count": char_count,
                "parser_used": method_used,
                "ocr_attempted": ocr_attempted,
                "ocr_used": ocr_used,
                "ocr_error": ocr_error
            }
        }

    def _extract_docx(self, path: Path, filename: str) -> Dict[str, Any]:
        """Extract readable text from DOCX using python-docx with zipfile fallback."""
        text = ""
        paragraphs_count = 0
        method_used = "python-docx"

        # Attempt 1: python-docx
        try:
            import docx
            doc = docx.Document(str(path))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            paragraphs_count = len(paragraphs)
            text = "\n".join(paragraphs[:20])

            # Also check table contents if paragraphs are brief
            if len(text) < 200 and doc.tables:
                table_texts = []
                for table in doc.tables[:3]:
                    for row in table.rows[:4]:
                        row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_vals:
                            table_texts.append(" | ".join(row_vals))
                if table_texts:
                    text += "\n" + "\n".join(table_texts)
        except Exception:
            method_used = "zipfile_xml"
            # Attempt 2: built-in zipfile parsing of word/document.xml
            try:
                with zipfile.ZipFile(str(path), "r") as zf:
                    if "word/document.xml" in zf.namelist():
                        xml_content = zf.read("word/document.xml")
                        root = ET.fromstring(xml_content)
                        w_t_texts = [
                            elem.text for elem in root.iter()
                            if elem.tag.endswith("t") and elem.text
                        ]
                        text = " ".join(w_t_texts)
            except Exception:
                text = ""

        clean_text = self._clean_text(text)
        char_count = len(clean_text)
        has_content = char_count >= 15

        logger.info(
            f"[Diagnostic:Extraction] DOCX '{filename}' ({path.stat().st_size} bytes): "
            f"paragraphs={paragraphs_count}, text_extracted={has_content}, char_count={char_count}, parser='{method_used}'"
        )

        return {
            "has_content": has_content,
            "content_snippet": clean_text[:MAX_SNIPPET_LENGTH] if has_content else "",
            "content_type": "docx_text",
            "extracted_metadata": {
                "paragraphs_count": paragraphs_count,
                "is_empty_or_minimal": not has_content,
                "char_count": char_count
            }
        }

    def _extract_pptx(self, path: Path, filename: str) -> Dict[str, Any]:
        """Extract slide text from PowerPoint .pptx archives."""
        text = ""
        slide_count = 0
        try:
            with zipfile.ZipFile(str(path), "r") as zf:
                slide_files = [n for n in zf.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
                slide_count = len(slide_files)
                extracted_lines = []
                for s_file in sorted(slide_files)[:5]:
                    xml_data = zf.read(s_file)
                    root = ET.fromstring(xml_data)
                    slide_texts = [
                        elem.text.strip() for elem in root.iter()
                        if elem.tag.endswith("t") and elem.text and elem.text.strip()
                    ]
                    if slide_texts:
                        extracted_lines.append(" ".join(slide_texts))
                text = "\n".join(extracted_lines)
        except Exception:
            text = ""

        clean_text = self._clean_text(text)
        has_content = len(clean_text) >= 15

        logger.info(
            f"[Diagnostic:Extraction] PPTX '{filename}' ({path.stat().st_size} bytes): "
            f"slides={slide_count}, text_extracted={has_content}, char_count={len(clean_text)}"
        )

        return {
            "has_content": has_content,
            "content_snippet": clean_text[:MAX_SNIPPET_LENGTH] if has_content else "",
            "content_type": "pptx_text",
            "extracted_metadata": {
                "slide_count": slide_count,
                "is_empty_or_minimal": not has_content,
                "char_count": len(clean_text)
            }
        }

    def _extract_csv(self, path: Path, filename: str) -> Dict[str, Any]:
        """Extract column headers and sample data rows from CSV."""
        headers: List[str] = []
        sample_rows: List[List[str]] = []

        try:
            with open(path, "rb") as f:
                raw_bytes = f.read(16384)

            text_data = ""
            for enc in ("utf-8-sig", "utf-8", "latin-1"):
                try:
                    text_data = raw_bytes.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue

            if text_data:
                delimiter = ","
                try:
                    sniffer = csv.Sniffer()
                    dialect = sniffer.sniff(text_data[:2048])
                    delimiter = dialect.delimiter
                except Exception:
                    for d in (",", ";", "\t", "|"):
                        if d in text_data:
                            delimiter = d
                            break

                reader = csv.reader(io.StringIO(text_data), delimiter=delimiter)
                all_rows = []
                for row in reader:
                    cleaned_row = [col.strip() for col in row if col.strip()]
                    if cleaned_row:
                        all_rows.append(cleaned_row)
                    if len(all_rows) >= (MAX_PREVIEW_ROWS + 1):
                        break

                if all_rows:
                    headers = all_rows[0]
                    sample_rows = all_rows[1:MAX_PREVIEW_ROWS + 1]

        except Exception as e:
            logger.debug(f"CSV read error: {e}")

        lines = []
        if headers:
            lines.append(f"Columns: [{', '.join(headers[:8])}]")
        if sample_rows:
            lines.append("Sample rows:")
            for idx, row in enumerate(sample_rows[:3], 1):
                row_str = ", ".join(row[:8])
                lines.append(f"  Row {idx}: {row_str}")

        snippet = "\n".join(lines)
        has_content = bool(headers or sample_rows)

        logger.info(
            f"[Diagnostic:Extraction] CSV '{filename}' ({path.stat().st_size} bytes): "
            f"headers={len(headers)}, sample_rows={len(sample_rows)}, text_extracted={has_content}"
        )

        return {
            "has_content": has_content,
            "content_snippet": snippet[:MAX_SNIPPET_LENGTH],
            "content_type": "csv_data",
            "extracted_metadata": {
                "headers": headers[:10],
                "header_count": len(headers),
                "sample_row_count": len(sample_rows),
                "is_empty_or_minimal": not has_content
            }
        }

    def _extract_xlsx(self, path: Path, filename: str) -> Dict[str, Any]:
        """Extract sheet names, column headers, and top sample rows from XLSX."""
        sheet_names: List[str] = []
        headers: List[str] = []
        sample_rows: List[List[str]] = []

        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            sheet_names = wb.sheetnames
            ws = wb.active or wb[sheet_names[0]]

            row_iter = ws.iter_rows(values_only=True)
            rows = []
            for row in row_iter:
                vals = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if vals:
                    rows.append(vals)
                if len(rows) >= (MAX_PREVIEW_ROWS + 1):
                    break
            wb.close()

            if rows:
                headers = rows[0]
                sample_rows = rows[1:MAX_PREVIEW_ROWS + 1]
        except Exception:
            try:
                with zipfile.ZipFile(str(path), "r") as zf:
                    if "xl/sharedStrings.xml" in zf.namelist():
                        xml_data = zf.read("xl/sharedStrings.xml")
                        root = ET.fromstring(xml_data)
                        strings = [elem.text for elem in root.iter() if elem.tag.endswith("t") and elem.text]
                        if strings:
                            headers = strings[:8]
                            sample_rows = [strings[8:16]] if len(strings) > 8 else []
            except Exception:
                pass

        lines = []
        if sheet_names:
            lines.append(f"Sheets: {', '.join(sheet_names[:4])}")
        if headers:
            lines.append(f"Columns: [{', '.join(headers[:8])}]")
        if sample_rows:
            lines.append("Sample rows:")
            for idx, row in enumerate(sample_rows[:3], 1):
                row_str = ", ".join(row[:8])
                lines.append(f"  Row {idx}: {row_str}")

        snippet = "\n".join(lines)
        has_content = bool(headers or sample_rows)

        logger.info(
            f"[Diagnostic:Extraction] XLSX '{filename}' ({path.stat().st_size} bytes): "
            f"sheets={len(sheet_names)}, headers={len(headers)}, text_extracted={has_content}"
        )

        return {
            "has_content": has_content,
            "content_snippet": snippet[:MAX_SNIPPET_LENGTH],
            "content_type": "xlsx_data",
            "extracted_metadata": {
                "sheet_names": sheet_names[:5],
                "headers": headers[:10],
                "is_empty_or_minimal": not has_content
            }
        }

    def _extract_image_metadata(self, path: Path, filename: str) -> Dict[str, Any]:
        """
        Extract visual text from image files (JPEG, PNG, WebP, etc.) using OCR.
        Falls back to image dimensions and EXIF metadata only when OCR finds no text.
        Never treats image metadata as a substitute for visual document text.
        """
        img_format = ""
        width = 0
        height = 0
        mode = ""
        exif_details = []

        try:
            from PIL import Image, ExifTags
            with Image.open(str(path)) as img:
                img_format = img.format or path.suffix.upper().lstrip(".")
                width, height = img.size
                mode = img.mode

                exif_data = img.getexif()
                if exif_data:
                    for tag_id, value in exif_data.items():
                        tag_name = ExifTags.TAGS.get(tag_id, str(tag_id))
                        if tag_name in ("DateTimeOriginal", "DateTime", "Make", "Model", "Software", "ImageDescription"):
                            clean_val = str(value).strip()
                            if clean_val:
                                exif_details.append(f"{tag_name}: {clean_val}")
        except Exception:
            pass

        # 1. Primary Strategy: Run OCR directly on image content
        ocr_attempted = False
        ocr_used = False
        ocr_error = None
        ocr_text = ""

        if ocr_service.is_available():
            ocr_attempted = True
            ocr_res = ocr_service.ocr_image(path, filename=filename)
            if ocr_res.get("success") and ocr_res.get("text"):
                clean_ocr = self._clean_text(ocr_res["text"])
                if len(clean_ocr) >= 15:
                    ocr_text = clean_ocr
                    ocr_used = True
            else:
                ocr_error = ocr_res.get("error")
        else:
            ocr_attempted = True
            ocr_error = "OCR engine not available"

        # Check if OCR extracted meaningful document / certificate content
        is_document_text = False
        if ocr_used and ocr_text:
            t_low = ocr_text.lower()
            doc_cues = (
                "certificate", "certify", "certifies", "awarded to", "completion",
                "invoice", "receipt", "total due", "amount due", "bill to",
                "tax", "statement", "curriculum vitae", "resume", "experience",
                "education", "course", "training", "program"
            )
            is_document_text = any(cue in t_low for cue in doc_cues) or len(ocr_text) >= 60

        if is_document_text:
            logger.info(
                f"[Diagnostic:Extraction] Image '{filename}' ({path.stat().st_size} bytes): "
                f"format={img_format}, res={width}x{height}, OCR extracted {len(ocr_text)} chars of document text"
            )
            return {
                "has_content": True,
                "content_snippet": ocr_text[:MAX_SNIPPET_LENGTH],
                "content_type": "image_ocr",
                "extracted_metadata": {
                    "format": img_format,
                    "width": width,
                    "height": height,
                    "char_count": len(ocr_text),
                    "parser_used": "ocr_tesseract",
                    "ocr_attempted": True,
                    "ocr_used": True,
                    "is_empty_or_minimal": False
                }
            }

        # 2. Secondary Strategy: Visual Image Classification
        # When OCR cannot identify document content, use the vision classifier
        vision_attempted = False
        vision_used = False

        if vision_service.is_available():
            vision_attempted = True
            vis_res = vision_service.classify_image(path, filename=filename)
            if vis_res.get("success") and vis_res.get("concept"):
                concept = vis_res["concept"]
                category = vis_res["category"]
                confidence = vis_res["confidence_level"]
                score = vis_res["confidence_score"]

                snippet = f"Visual classification: identified {concept} ({category}). Confidence: {confidence}."
                logger.info(
                    f"[Diagnostic:Extraction] Image '{filename}' ({path.stat().st_size} bytes): "
                    f"format={img_format}, visual_concept='{concept}', category='{category}', score={score:.2%}"
                )
                return {
                    "has_content": True,
                    "content_snippet": snippet,
                    "content_type": "image_visual",
                    "extracted_metadata": {
                        "format": img_format,
                        "width": width,
                        "height": height,
                        "visual_concept": concept,
                        "visual_category": category,
                        "visual_folder": vis_res.get("folder", "Images"),
                        "visual_confidence": confidence,
                        "confidence_score": score,
                        "raw_label": vis_res.get("raw_label", ""),
                        "ocr_attempted": ocr_attempted,
                        "ocr_used": False,
                        "vision_attempted": True,
                        "vision_used": True,
                        "is_empty_or_minimal": False
                    }
                }

        # 3. Fallback: Structural metadata for images without readable text or recognizable visual object
        # (ordinary photos, unclear images, abstract textures)
        heuristic = ""
        if (width, height) in ((1920, 1080), (2560, 1440), (1366, 768), (1440, 900), (3840, 2160)):
            heuristic = "Standard screen resolution (likely UI screenshot or graphic)."
        elif any(make_kw in " ".join(exif_details).lower() for make_kw in ("apple", "sony", "canon", "nikon", "samsung", "google")):
            heuristic = "Camera photograph."

        parts = [f"Image ({img_format}, {width}x{height}, color mode {mode})."]
        if exif_details:
            parts.append(f"EXIF: {'; '.join(exif_details[:3])}.")
        if heuristic:
            parts.append(heuristic)

        snippet = " ".join(parts)

        logger.info(
            f"[Diagnostic:Extraction] Image '{filename}' ({path.stat().st_size} bytes): "
            f"format={img_format}, res={width}x{height}, ocr_document_found=False, vision_classified=False"
        )

        return {
            "has_content": False,
            "content_snippet": snippet,
            "content_type": "image_metadata",
            "extracted_metadata": {
                "format": img_format,
                "width": width,
                "height": height,
                "is_empty_or_minimal": True,
                "is_ambiguous": True,
                "ocr_attempted": ocr_attempted,
                "ocr_used": False,
                "vision_attempted": vision_attempted,
                "vision_used": False
            }
        }

    def _extract_text_or_code(self, path: Path, filename: str, ext: str) -> Dict[str, Any]:
        """Extract preview from plain text, Markdown, scripts, or config files."""
        try:
            with open(path, "rb") as f:
                raw_bytes = f.read(2048)

            text = ""
            for enc in ("utf-8", "utf-8-sig", "latin-1"):
                try:
                    text = raw_bytes.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue

            clean = self._clean_text(text)
            has_content = len(clean.strip()) >= 10

            logger.info(
                f"[Diagnostic:Extraction] Text/Code '{filename}' ({path.stat().st_size} bytes): "
                f"ext='{ext}', text_extracted={has_content}, char_count={len(clean)}"
            )

            return {
                "has_content": has_content,
                "content_snippet": clean[:MAX_SNIPPET_LENGTH] if has_content else "",
                "content_type": "text_preview",
                "extracted_metadata": {
                    "extension": ext,
                    "is_empty_or_minimal": not has_content,
                    "char_count": len(clean)
                }
            }
        except Exception:
            return {
                "has_content": False,
                "content_snippet": "",
                "content_type": "text_error",
                "extracted_metadata": {"is_empty_or_minimal": True}
            }

    def _extract_generic_binary_info(
        self,
        path: Path,
        filename: str,
        ext: str,
        mime: str
    ) -> Dict[str, Any]:
        """Non-text binary files: note format and size without transmitting raw bytes."""
        size = path.stat().st_size
        logger.info(f"[Diagnostic:Extraction] Binary file '{filename}' ({size} bytes): format='{ext or 'bin'}'")
        return {
            "has_content": False,
            "content_snippet": "",
            "content_type": "binary_info",
            "extracted_metadata": {
                "size_bytes": size,
                "is_empty_or_minimal": True
            }
        }

    def _clean_text(self, text: str) -> str:
        """Collapse redundant whitespace, expand unicode ligatures, re-join hyphenated line breaks, and remove non-printable characters."""
        if not text:
            return ""
        return ocr_service.normalize_text(text)


content_extractor = ContentExtractor()
